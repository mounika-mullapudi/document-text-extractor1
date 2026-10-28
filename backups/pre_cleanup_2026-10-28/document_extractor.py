
from pathlib import Path
import csv
import zipfile

import fitz
import easyocr
from docx import Document
from openpyxl import load_workbook
from pdf_combined_extraction import extract_combined_pdf


SUPPORTED_FORMATS = [".pdf", ".docx", ".xlsx", ".csv"]
IMAGE_FORMATS = [".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"]

# Initialize OCR reader once
ocr_reader = easyocr.Reader(["en"], gpu=False)


def extract_image_text(file_path):
    """Extract image text and group words by visual lines."""
    try:
        results = ocr_reader.readtext(
            str(file_path),
            detail=1,
            paragraph=False
        )

        items = []

        for box, text, confidence in results:
            text = text.strip()
            if not text:
                continue

            left = min(point[0] for point in box)
            top = min(point[1] for point in box)
            bottom = max(point[1] for point in box)
            center_y = (top + bottom) / 2

            items.append({
                "text": text,
                "left": left,
                "top": top,
                "bottom": bottom,
                "center_y": center_y,
                "height": bottom - top,
            })

        # Sort detected boxes from top to bottom
        items.sort(key=lambda item: (item["center_y"], item["left"]))

        lines = []

        for item in items:
            matching_line = None

            for line in lines:
                average_y = sum(
                    word["center_y"] for word in line
                ) / len(line)

                average_height = sum(
                    word["height"] for word in line
                ) / len(line)

                tolerance = max(10, average_height * 0.5)

                if abs(item["center_y"] - average_y) <= tolerance:
                    matching_line = line
                    break

            if matching_line is None:
                lines.append([item])
            else:
                matching_line.append(item)

        # Sort words left to right within each line
        output_lines = []

        for line in lines:
            line.sort(key=lambda item: item["left"])
            output_lines.append(
                " ".join(item["text"] for item in line)
            )

        return "\n".join(output_lines).strip()

    except Exception as error:
        raise RuntimeError(f"Image OCR failed: {error}") from error


def extract_pdf(file_path):
    """Extract and format PDF text, tables, and scanned-page OCR."""
    return format_pdf_extraction(extract_pdf_data(file_path))


def format_pdf_extraction(extraction):
    """Format the article's page results for display and TXT download."""
    output = []

    for page_name, page in extraction.items():
        if page_name == "warnings":
            continue

        output.append(f"\n--- {page_name.replace('_', ' ')} ---")
        output.extend(page["page_text"])

        for table_text in page["text_from_tables"]:
            output.append(f"\n[Table]\n{table_text}")

        output.extend(page["text_from_images"])

    return "\n".join(output).strip()


def extract_pdf_data(file_path):
    """Return the structured combined PDF extraction for the interface."""
    return extract_combined_pdf(file_path)


def extract_docx(file_path):
    """Extract paragraphs, tables, headers, and footers from DOCX."""
    doc = Document(file_path)
    output = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            output.append(paragraph.text.strip())

    for table_number, table in enumerate(doc.tables, start=1):
        output.append(f"\n[Table {table_number}]")

        for row in table.rows:
            output.append(
                " | ".join(cell.text.strip() for cell in row.cells)
            )

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                output.append(f"[Header] {paragraph.text.strip()}")

        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                output.append(f"[Footer] {paragraph.text.strip()}")

    return "\n".join(output).strip()


def extract_xlsx(file_path):
    """Extract non-empty cell values from all worksheets."""
    workbook = load_workbook(file_path, data_only=True)
    output = []

    try:
        for sheet in workbook.worksheets:
            output.append(f"\n--- Sheet: {sheet.title} ---")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    output.append(" | ".join(values))
    finally:
        workbook.close()

    return "\n".join(output).strip()


def extract_csv(file_path):
    """Extract CSV data with common delimiter detection."""
    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:
        sample = file.read(8192)
        file.seek(0)

        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(file, dialect)

        return "\n".join(
            " | ".join(cell.strip() for cell in row)
            for row in reader
            if any(cell.strip() for cell in row)
        ).strip()


def extract_text(file_path):
    """Select the appropriate extractor by file extension."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)
    elif extension == ".docx":
        return extract_docx(path)
    elif extension == ".xlsx":
        return extract_xlsx(path)
    elif extension == ".csv":
        return extract_csv(path)
    elif extension in IMAGE_FORMATS:
        return extract_image_text(path)

    raise ValueError(f"Unsupported file format: {extension}")


def extract_document_images(file_path):
    """Extract embedded images and PDF page previews."""
    path = Path(file_path)
    extension = path.suffix.lower()
    images = []

    if extension == ".pdf":
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                for image_number, info in enumerate(
                    page.get_images(full=True), start=1
                ):
                    try:
                        extracted = pdf.extract_image(info[0])

                        images.append({
                            "name": f"page_{page_number}_image_{image_number}",
                            "data": extracted["image"],
                            "extension": extracted["ext"]
                        })
                    except Exception:
                        pass

                pixmap = page.get_pixmap(
                    matrix=fitz.Matrix(1.5, 1.5),
                    alpha=False
                )

                images.append({
                    "name": f"page_{page_number}_preview",
                    "data": pixmap.tobytes("png"),
                    "extension": "png"
                })

    elif extension == ".docx":
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    images.append({
                        "name": Path(name).name,
                        "data": archive.read(name),
                        "extension": Path(name).suffix.lstrip(".") or "bin"
                    })

    elif extension == ".xlsx":
        workbook = load_workbook(path)

        try:
            for sheet in workbook.worksheets:
                for number, image in enumerate(
                    getattr(sheet, "_images", []), start=1
                ):
                    try:
                        images.append({
                            "name": f"{sheet.title}_image_{number}",
                            "data": image._data(),
                            "extension": getattr(image, "format", None) or "png"
                        })
                    except Exception:
                        pass
        finally:
            workbook.close()

    elif extension in IMAGE_FORMATS:
        images.append({
            "name": path.stem,
            "data": path.read_bytes(),
            "extension": extension.lstrip(".")
        })

    return images
